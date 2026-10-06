from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest import mock

from map_builder.thematic_wgi_ingest import (
    DEFAULT_WGI_SOURCE_CACHE_PATH,
    DIMENSION_SHEETS,
    DIMENSION_TO_METRIC_ID,
    WGI_AUDIT_RELATIVE_PATH,
    WGI_COMPOSITE_METRIC_ID,
    WGI_GOVERNMENT_EFFECTIVENESS_METRIC_ID,
    WGI_LAYER_ID,
    WGI_MANIFEST_RELATIVE_PATH,
    WGI_METRICS_RELATIVE_PATH,
    WGI_RECIPE_RELATIVE_PATH,
    WGI_RULE_OF_LAW_METRIC_ID,
    WGI_RUNTIME_METRIC_IDS,
    WGI_RUNTIME_SCENARIO_IDS,
    WGI_RUNTIME_DATA_VERSION,
    WGI_RUNTIME_METHOD,
    build_wgi_historical_reference_policy,
    build_wgi_real_source_payloads,
    source_code_to_join_key,
)
from map_builder.thematic_layer_contracts import (
    validate_thematic_admin_metrics,
    validate_thematic_build_audit,
    validate_thematic_layer_manifest,
)


REPO_ROOT = Path(__file__).resolve().parents[1]
FIXTURE_PATH = REPO_ROOT / "tests" / "fixtures" / "thematic_wgi_2024_minimal.csv"


def _feature_by_join_key(metrics_payload: dict, join_key: str) -> dict:
    for feature in metrics_payload["features"]:
        if feature["join_key"] == join_key:
            return feature
    raise AssertionError(f"feature not found: {join_key}")


class ThematicWgiSourceIngestTest(unittest.TestCase):
    @unittest.skipUnless(DEFAULT_WGI_SOURCE_CACHE_PATH.is_file(), "Official WGI workbook cache is not available")
    def test_official_workbook_passes_through_all_six_dimensions_and_uncertainty(self) -> None:
        from openpyxl import load_workbook

        payloads = build_wgi_real_source_payloads(DEFAULT_WGI_SOURCE_CACHE_PATH)
        features = {feature["join_key"]: feature for feature in payloads[WGI_METRICS_RELATIVE_PATH]["features"]}
        audit = payloads[WGI_AUDIT_RELATIVE_PATH]["coverage_summary"]["dimensions"]
        workbook = load_workbook(DEFAULT_WGI_SOURCE_CACHE_PATH, read_only=True, data_only=True)
        source_join_keys = set()
        try:
            for sheet_name, dimension in DIMENSION_SHEETS.items():
                rows = workbook[sheet_name].iter_rows(values_only=True)
                headers = next(rows)
                selected_rows = mapped_rows = observed_rows = 0
                for cells in rows:
                    source = dict(zip(headers, cells))
                    if source["Year"] != 2024:
                        continue
                    selected_rows += 1
                    join_key = source_code_to_join_key(source["Economy (code)"])
                    if join_key is None:
                        continue
                    mapped_rows += 1
                    source_join_keys.add(join_key)
                    metric = features[join_key]["values"][DIMENSION_TO_METRIC_ID[dimension]]
                    raw_score = source["Governance score (0-100)"]
                    expected_score = round(raw_score, 6) if isinstance(raw_score, (int, float)) else None
                    self.assertEqual(metric["raw_value"], expected_score)
                    self.assertEqual(metric["normalized_value"], expected_score)
                    self.assertEqual(metric["source_status"], "observed" if expected_score is not None else "source_gap")
                    observed_rows += expected_score is not None
                    expected_uncertainty = {
                        "number_of_sources": source["Number of sources"],
                        "score_standard_error": source["Standard error (gov. score)"],
                        "score_confidence_interval_90": {
                            "lower": source["Lower threshold (90% conf. int. score)"],
                            "upper": source["Upper threshold (90% conf. int. score)"],
                        },
                        "estimate": source["Governance estimate (approx. -2.5 to +2.5)"],
                        "estimate_standard_error": source["Standard error (estimate)"],
                        "estimate_confidence_interval_90": {
                            "lower": source["Lower threshold (90% conf. int. estimate)"],
                            "upper": source["Upper threshold (90% conf. int. estimate)"],
                        },
                    }
                    def rounded(value):
                        if isinstance(value, dict):
                            return {key: rounded(item) for key, item in value.items()}
                        return round(value, 6) if isinstance(value, (int, float)) else None
                    self.assertEqual(metric["uncertainty"], rounded(expected_uncertainty))
                self.assertEqual(audit[dimension]["selected_source_rows"], selected_rows)
                self.assertEqual(audit[dimension]["source_rows_mapped"], mapped_rows)
                self.assertEqual(audit[dimension]["observed"], observed_rows)
        finally:
            workbook.close()
        self.assertEqual(set(features), source_join_keys)
        self.assertEqual(len(WGI_RUNTIME_METRIC_IDS), 6)
        self.assertNotIn(WGI_COMPOSITE_METRIC_ID, WGI_RUNTIME_METRIC_IDS)
        for feature in features.values():
            values = feature["values"]
            ge = values[WGI_GOVERNMENT_EFFECTIVENESS_METRIC_ID]["raw_value"]
            rl = values[WGI_RULE_OF_LAW_METRIC_ID]["raw_value"]
            expected = round((ge + rl) / 2, 6) if ge is not None and rl is not None else None
            self.assertEqual(values[WGI_COMPOSITE_METRIC_ID]["raw_value"], expected)

    def test_six_dimensions_use_source_union_and_audit_each_dimension(self) -> None:
        rows = ["Economy (name),Economy (code),Year,Governance dimension,Governance score (0-100),Number of sources"]
        for index, (sheet, dimension) in enumerate(DIMENSION_SHEETS.items()):
            rows.append(f"United States,USA,2024,{sheet},{70 + index},10")
        rows.extend([
            "Canada,CAN,2024,va,64,8",
            "Mexico,MEX,2024,pv,,7",
            "World,WLD,2024,rq,60,9",
            "Unknown,ZZZ,2024,cc,50,6",
        ])
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "six-dimensions.csv"
            path.write_text("\n".join(rows), encoding="utf-8")
            payloads = build_wgi_real_source_payloads(path)
        metrics = payloads[WGI_METRICS_RELATIVE_PATH]
        audit = payloads[WGI_AUDIT_RELATIVE_PATH]["coverage_summary"]["dimensions"]
        self.assertEqual({feature["join_key"] for feature in metrics["features"]}, {"USA", "CAN", "MEX"})
        self.assertEqual(_feature_by_join_key(metrics, "USA")["coverage_status"], "complete")
        canada = _feature_by_join_key(metrics, "CAN")
        self.assertEqual(canada["coverage_status"], "partial")
        self.assertEqual(canada["values"][WGI_COMPOSITE_METRIC_ID]["source_status"], "source_gap")
        mexico = _feature_by_join_key(metrics, "MEX")
        political_stability = mexico["values"][DIMENSION_TO_METRIC_ID["political_stability"]]
        self.assertEqual(political_stability["source_status"], "source_gap")
        self.assertEqual(political_stability["uncertainty"]["number_of_sources"], 7)
        self.assertEqual(mexico["coverage_status"], "missing")
        self.assertEqual(audit["voice_and_accountability"]["source_rows_mapped"], 2)
        self.assertEqual(audit["political_stability"]["source_rows_with_missing_score"], 1)
        self.assertEqual(audit["regulatory_quality"]["source_rows_dropped_aggregate"], 1)
        self.assertEqual(audit["control_of_corruption"]["source_rows_unmatched"], 1)
        for dimension, counts in audit.items():
            self.assertEqual(counts["selected_source_rows"], counts["source_rows_mapped"] + counts["source_rows_unmatched"] + counts["source_rows_dropped_aggregate"])
            self.assertEqual(counts["observed"] + counts["source_gap"], 3)
            self.assertEqual(counts["source_rows_mapped"] + counts["missing_source_rows"], 3)

    def _build_fixture_payloads(self) -> dict[str, dict]:
        with mock.patch("urllib.request.urlopen") as urlopen:
            payloads = build_wgi_real_source_payloads(
                FIXTURE_PATH,
                generated_at="2026-06-22T00:00:00Z",
                accessed_at="2026-06-22T00:00:00Z",
            )
        urlopen.assert_not_called()
        return payloads

    def test_real_source_request_requires_existing_wgi_cache_without_network(self) -> None:
        missing_path = REPO_ROOT / ".runtime" / "source-cache" / "thematic" / "wgi" / "missing.xlsx"

        with mock.patch("urllib.request.urlopen") as urlopen:
            with self.assertRaises(FileNotFoundError) as raised:
                build_wgi_real_source_payloads(missing_path)

        self.assertIn(str(missing_path.resolve()), str(raised.exception))
        urlopen.assert_not_called()

    def test_fixture_wgi_ingest_builds_contract_valid_payloads_without_network(self) -> None:
        payloads = self._build_fixture_payloads()
        manifest = payloads[WGI_MANIFEST_RELATIVE_PATH]
        metrics = payloads[WGI_METRICS_RELATIVE_PATH]
        audit = payloads[WGI_AUDIT_RELATIVE_PATH]
        recipe = payloads[WGI_RECIPE_RELATIVE_PATH]

        self.assertEqual(validate_thematic_layer_manifest(manifest), [])
        self.assertEqual(validate_thematic_admin_metrics(metrics), [])
        self.assertEqual(validate_thematic_build_audit(audit), [])
        self.assertEqual(manifest["layer_id"], WGI_LAYER_ID)
        self.assertEqual(manifest["source_policy"], "real_source_cache_only")
        self.assertEqual(manifest["coverage_scope"]["join_key_type"], "iso_a3")
        self.assertEqual(recipe["download_policy"]["network_allowed"], False)
        self.assertEqual(manifest["runtime_consumer"]["supported_metrics"], list(WGI_RUNTIME_METRIC_IDS))
        self.assertEqual(recipe["runtime_selection"]["supported_metrics"], list(WGI_RUNTIME_METRIC_IDS))
        for selection in (manifest["runtime_consumer"], recipe["runtime_selection"]):
            self.assertEqual(selection["supported_scenarios"], list(WGI_RUNTIME_SCENARIO_IDS))
            self.assertEqual(selection["data_version"], WGI_RUNTIME_DATA_VERSION)
            self.assertEqual(selection["method"], WGI_RUNTIME_METHOD)
            self.assertEqual(selection["historical_reference_policy"], build_wgi_historical_reference_policy())
            self.assertEqual(selection["historical_reference_policy"]["reference_year"], 2024)
            self.assertFalse(selection["historical_reference_policy"]["historical_measurement"])

    def test_wgi_ingest_builds_government_effectiveness_rule_of_law_and_composite(self) -> None:
        metrics = self._build_fixture_payloads()[WGI_METRICS_RELATIVE_PATH]
        usa = _feature_by_join_key(metrics, "USA")
        values = usa["values"]

        self.assertEqual(values[WGI_GOVERNMENT_EFFECTIVENESS_METRIC_ID]["normalized_value"], 80.0)
        self.assertEqual(values[WGI_RULE_OF_LAW_METRIC_ID]["normalized_value"], 72.0)
        self.assertEqual(values[WGI_COMPOSITE_METRIC_ID]["normalized_value"], 76.0)
        self.assertEqual(values[WGI_COMPOSITE_METRIC_ID]["source_status"], "observed")
        self.assertEqual(values[WGI_COMPOSITE_METRIC_ID]["year"], 2024)
        ge_uncertainty = values[WGI_GOVERNMENT_EFFECTIVENESS_METRIC_ID]["uncertainty"]
        self.assertEqual(ge_uncertainty["number_of_sources"], 12)
        self.assertEqual(ge_uncertainty["score_standard_error"], 4.2)
        self.assertEqual(ge_uncertainty["score_confidence_interval_90"], {"lower": 72.0, "upper": 88.0})
        self.assertEqual(ge_uncertainty["estimate"], 1.2)
        self.assertEqual(ge_uncertainty["estimate_standard_error"], 0.15)
        self.assertEqual(ge_uncertainty["estimate_confidence_interval_90"], {"lower": 0.9, "upper": 1.5})
        composite_uncertainty = values[WGI_COMPOSITE_METRIC_ID]["uncertainty"]
        self.assertEqual(composite_uncertainty["method"], "not_computed")
        self.assertIn("not inferred", composite_uncertainty["reason"])

    def test_wgi_join_uses_explicit_codes_without_fuzzy_name_matching(self) -> None:
        metrics = self._build_fixture_payloads()[WGI_METRICS_RELATIVE_PATH]
        join_keys = {feature["join_key"] for feature in metrics["features"]}

        self.assertEqual(source_code_to_join_key("ADO"), "AND")
        self.assertEqual(source_code_to_join_key("XKX"), "XKX")
        self.assertIsNone(source_code_to_join_key("CHI"))
        self.assertIsNone(source_code_to_join_key("ZZZ"))
        self.assertIn("USA", join_keys)
        self.assertIn("AND", join_keys)
        self.assertIn("XKX", join_keys)
        self.assertNotIn("CHI", join_keys)
        self.assertNotIn("ZZZ", join_keys)

    def test_wgi_ingest_drops_aggregate_rows_into_audit(self) -> None:
        payloads = self._build_fixture_payloads()
        metrics = payloads[WGI_METRICS_RELATIVE_PATH]
        audit = payloads[WGI_AUDIT_RELATIVE_PATH]

        join_keys = {feature["join_key"] for feature in metrics["features"]}
        dropped_codes = {row["source_code"] for row in audit["dropped_aggregate_rows"]}

        self.assertNotIn("WLD", join_keys)
        self.assertNotIn("HIC", join_keys)
        self.assertEqual(dropped_codes, {"WLD", "HIC"})

    def test_wgi_unmatched_rows_are_audited(self) -> None:
        audit = self._build_fixture_payloads()[WGI_AUDIT_RELATIVE_PATH]
        unmatched = audit["unmatched_source_rows"]
        unmatched_by_code = {row["source_code"]: row for row in unmatched}

        self.assertEqual(set(unmatched_by_code), {"CHI", "ZZZ"})
        self.assertEqual(unmatched_by_code["CHI"]["reason"], "unmatched_join_key")
        self.assertEqual(unmatched_by_code["CHI"]["dimension"], "government_effectiveness")
        self.assertEqual(unmatched_by_code["ZZZ"]["reason"], "unmatched_join_key")
        self.assertEqual(unmatched_by_code["ZZZ"]["dimension"], "rule_of_law")

    def test_wgi_missing_one_metric_sets_composite_partial_source_gap(self) -> None:
        metrics = self._build_fixture_payloads()[WGI_METRICS_RELATIVE_PATH]
        brazil = _feature_by_join_key(metrics, "BRA")
        values = brazil["values"]

        self.assertEqual(brazil["coverage_status"], "partial")
        self.assertEqual(values[WGI_GOVERNMENT_EFFECTIVENESS_METRIC_ID]["source_status"], "observed")
        self.assertIsNone(values[WGI_RULE_OF_LAW_METRIC_ID]["raw_value"])
        self.assertEqual(values[WGI_RULE_OF_LAW_METRIC_ID]["source_status"], "source_gap")
        self.assertIsNone(values[WGI_COMPOSITE_METRIC_ID]["normalized_value"])
        self.assertEqual(values[WGI_COMPOSITE_METRIC_ID]["source_status"], "partial_source_gap")

    def test_wgi_missing_both_metrics_sets_source_gap(self) -> None:
        metrics = self._build_fixture_payloads()[WGI_METRICS_RELATIVE_PATH]
        south_africa = _feature_by_join_key(metrics, "ZAF")
        values = south_africa["values"]

        self.assertEqual(south_africa["coverage_status"], "missing")
        self.assertIsNone(values[WGI_GOVERNMENT_EFFECTIVENESS_METRIC_ID]["raw_value"])
        self.assertIsNone(values[WGI_RULE_OF_LAW_METRIC_ID]["raw_value"])
        self.assertIsNone(values[WGI_COMPOSITE_METRIC_ID]["raw_value"])
        self.assertEqual(values[WGI_COMPOSITE_METRIC_ID]["source_status"], "source_gap")

    def test_wgi_rejects_non_finite_and_out_of_range_scores_as_source_gap(self) -> None:
        rows = "\n".join(
            [
                "Economy (name),Economy (code),Year,Governance dimension,Governance score (0-100)",
                "United States,USA,2024,Government Effectiveness,NaN",
                "United States,USA,2024,Rule of Law,Infinity",
                "Canada,CAN,2024,Government Effectiveness,101",
                "Canada,CAN,2024,Rule of Law,-1",
                "Mexico,MEX,2024,Government Effectiveness,44",
                "Mexico,MEX,2024,Rule of Law,56",
            ]
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "wgi-invalid-scores.csv"
            source_path.write_text(rows, encoding="utf-8")
            metrics = build_wgi_real_source_payloads(source_path)[WGI_METRICS_RELATIVE_PATH]

        usa = _feature_by_join_key(metrics, "USA")
        canada = _feature_by_join_key(metrics, "CAN")
        mexico = _feature_by_join_key(metrics, "MEX")

        self.assertEqual(usa["coverage_status"], "missing")
        self.assertIsNone(usa["values"][WGI_GOVERNMENT_EFFECTIVENESS_METRIC_ID]["raw_value"])
        self.assertIsNone(usa["values"][WGI_RULE_OF_LAW_METRIC_ID]["raw_value"])
        self.assertEqual(canada["coverage_status"], "missing")
        self.assertIsNone(canada["values"][WGI_GOVERNMENT_EFFECTIVENESS_METRIC_ID]["raw_value"])
        self.assertIsNone(canada["values"][WGI_RULE_OF_LAW_METRIC_ID]["raw_value"])
        self.assertEqual(mexico["values"][WGI_COMPOSITE_METRIC_ID]["normalized_value"], 50.0)

    def test_wgi_number_of_sources_requires_non_negative_integer(self) -> None:
        rows = "\n".join(
            [
                "Economy (name),Economy (code),Year,Governance dimension,Governance score (0-100),Number of sources",
                "United States,USA,2024,Government Effectiveness,44,12.5",
                "United States,USA,2024,Rule of Law,56,7.0",
                "Canada,CAN,2024,Government Effectiveness,62,-1",
                "Canada,CAN,2024,Rule of Law,58,8",
            ]
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "wgi-source-counts.csv"
            source_path.write_text(rows, encoding="utf-8")
            metrics = build_wgi_real_source_payloads(source_path)[WGI_METRICS_RELATIVE_PATH]

        usa = _feature_by_join_key(metrics, "USA")
        canada = _feature_by_join_key(metrics, "CAN")

        self.assertIsNone(usa["values"][WGI_GOVERNMENT_EFFECTIVENESS_METRIC_ID]["uncertainty"]["number_of_sources"])
        self.assertEqual(usa["values"][WGI_RULE_OF_LAW_METRIC_ID]["uncertainty"]["number_of_sources"], 7)
        self.assertIsNone(canada["values"][WGI_GOVERNMENT_EFFECTIVENESS_METRIC_ID]["uncertainty"]["number_of_sources"])
        self.assertEqual(canada["values"][WGI_RULE_OF_LAW_METRIC_ID]["uncertainty"]["number_of_sources"], 8)


if __name__ == "__main__":
    unittest.main()
